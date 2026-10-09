"""Bulk doctor upload: read .xlsx/.xls/.csv, map columns, validate, match."""

from __future__ import annotations

import csv
import io
import re
from typing import Any

from .doctor_matcher import doctor_areas
from .mapping import norm_email, norm_name, norm_phone, norm_text
from .whatsapp_service import whatsapp_number

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000

FIELDS: list[tuple[str, list[str]]] = [
    ("doctor_name", ["doctor name", "name", "doctor", "dr name", "full name", "name of doctor"]),
    ("specialty", ["specialty", "speciality", "specialization", "specialisation", "designation"]),
    ("hospital", ["hospital", "institute", "hospital name", "institution", "organisation", "organization", "clinic"]),
    ("whatsapp_number", ["whatsapp number", "whatsapp", "whatsapp no", "mobile", "mobile number", "phone",
                         "contact number", "contact", "phone number", "mobile no"]),
    ("email", ["email", "email id", "e-mail", "email address", "mail id", "mail"]),
    ("department", ["department", "dept", "therapy area", "therapeutic area", "division"]),
    ("frequency", ["frequency", "contact frequency", "message frequency"]),
    ("must_see", ["must see", "must-see", "mustsee", "priority", "key doctor"]),
    ("city", ["city", "town", "location"]),
    ("state", ["state", "region"]),
    ("country", ["country"]),
    ("active", ["active", "status", "is active"]),
    ("notes", ["notes", "note", "remarks", "comments", "comment"]),
]
FIELD_NAMES = [f for f, _ in FIELDS]


class UploadError(Exception):
    pass


def _norm_header(h: Any) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(h or "").lower()).strip()


def read_table(filename: str, data: bytes) -> tuple[list[str], list[list[str]]]:
    if len(data) > MAX_BYTES:
        raise UploadError("The file is larger than 5 MB.")
    name = filename.lower()
    rows: list[list[Any]]
    if name.endswith(".csv"):
        text = None
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise UploadError("Could not read the CSV text encoding.")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.reader(io.StringIO(text), dialect))
    elif name.endswith(".xlsx"):
        if not data.startswith(b"PK"):
            raise UploadError("This is not a valid .xlsx file.")
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
    elif name.endswith(".xls"):
        if not data.startswith(b"\xd0\xcf\x11\xe0"):
            raise UploadError("This is not a valid .xls file.")
        import xlrd
        book = xlrd.open_workbook(file_contents=data)
        sh = book.sheet_by_index(0)
        rows = [sh.row_values(i) for i in range(sh.nrows)]
    else:
        raise UploadError("Upload a .xlsx, .xls or .csv file.")

    def cell(v: Any) -> str:
        if v is None:
            return ""
        if isinstance(v, float) and v.is_integer():
            return str(int(v))
        return str(v).strip()

    rows = [[cell(v) for v in r] for r in rows if r and any(cell(v) for v in r)]
    if not rows:
        raise UploadError("The file is empty.")
    # header = first row with at least two non-empty cells
    h = next((i for i, r in enumerate(rows[:10]) if sum(1 for c in r if c) >= 2), 0)
    header = rows[h]
    body = rows[h + 1:]
    if len(body) > MAX_ROWS:
        raise UploadError(f"The file has {len(body)} rows; the limit is {MAX_ROWS}.")
    width = len(header)
    body = [(r + [""] * width)[:width] for r in body]
    return header, body


def auto_mapping(header: list[str]) -> dict[str, int | None]:
    normed = [_norm_header(h) for h in header]
    mapping: dict[str, int | None] = {}
    used: set[int] = set()
    for field, synonyms in FIELDS:
        idx = None
        for syn in synonyms:
            for i, h in enumerate(normed):
                if i not in used and h == syn:
                    idx = i
                    break
            if idx is not None:
                break
        if idx is None:
            for syn in synonyms:
                for i, h in enumerate(normed):
                    if i not in used and len(syn) > 3 and syn in h:
                        idx = i
                        break
                if idx is not None:
                    break
        mapping[field] = idx
        if idx is not None:
            used.add(idx)
    return mapping


def _frequency(value: str) -> tuple[str | None, int | None]:
    v = value.strip().lower()
    if not v:
        return None, None
    if v.startswith("week") or v in ("7", "7 days"):
        return "WEEKLY", 7
    if "15" in v or "fortnight" in v:
        return "15_DAYS", 15
    if v.startswith("month") or v in ("30", "30 days"):
        return "MONTHLY", 30
    m = re.search(r"(\d+)", v)
    if m and 1 <= int(m.group(1)) <= 365:
        return "CUSTOM", int(m.group(1))
    return None, None


def _yes(value: str) -> bool | None:
    v = value.strip().lower()
    if not v:
        return None
    if v in ("yes", "y", "true", "1", "must see", "✓", "x"):
        return True
    if v in ("no", "n", "false", "0", "regular"):
        return False
    return None


def _department(value: str, specialty: str) -> tuple[str | None, str | None]:
    """App department + sub-department from the Department column, else the specialty."""
    v = f"{value} ".lower()
    s = specialty.lower()
    if "rare" in v or "genetic" in v:
        return "RARE_DISEASES", "Genetics" if "genetic" in (v + s) else None
    if "hema" in v or "haema" in v:
        return "NPP", "Hematology"
    if "onco" in v:
        return "NPP", "Oncology"
    if "npp" in v:
        return "NPP", ("Hematology" if re.search(r"hema|haema", s) else "Oncology" if "onc" in s else None)
    if re.search(r"genetic|rare|metabolic", s):
        return "RARE_DISEASES", "Genetics"
    if re.search(r"neuro", s):
        return "RARE_DISEASES", "Neuro"
    if re.search(r"hema|haema|bmt", s):
        return "NPP", "Hematology"
    if re.search(r"onc|cancer", s):
        return "NPP", "Oncology"
    return None, None


def validate(rows: list[list[str]], mapping: dict[str, int | None], existing: list[dict], allowed_departments: list[str],
             default_department: str | None) -> dict[str, Any]:
    def get(r: list[str], f: str) -> str:
        i = mapping.get(f)
        return r[i].strip() if i is not None and 0 <= i < len(r) else ""

    by_phone = {d["phone_norm"]: d for d in existing if d.get("phone_norm")}
    by_email = {d["email_norm"]: d for d in existing if d.get("email_norm")}
    by_name = {(d.get("name_norm"), norm_text(d.get("institute")) or ""): d for d in existing if d.get("name_norm")}

    out = []
    seen: dict[str, int] = {}
    summary = {k: 0 for k in ("total", "existing", "new", "oncology", "hematology", "rare_disease", "must_see",
                              "invalid_whatsapp", "missing_email", "duplicates", "errors")}
    for n, r in enumerate(rows):
        name = get(r, "doctor_name")
        if not name:
            continue
        summary["total"] += 1
        specialty = get(r, "specialty")
        issues: list[str] = []
        email_norm, email_issue = norm_email(get(r, "email"))
        phone_raw = get(r, "whatsapp_number")
        phone_norm, phone_issue = norm_phone(phone_raw)
        wa = whatsapp_number(phone_raw) if phone_norm else None
        if phone_raw and not wa:
            issues.append("Invalid WhatsApp number")
            summary["invalid_whatsapp"] += 1
        elif not phone_raw:
            issues.append("No WhatsApp number")
        if not email_norm:
            issues.append("Invalid email" if get(r, "email") else "Missing email")
            summary["missing_email"] += 1
        dept, sub = _department(get(r, "department"), specialty)
        if not dept:
            dept = default_department
            if not dept:
                issues.append("Department could not be determined")
        if dept and dept not in allowed_departments:
            issues.append("Belongs to a department you cannot access")
        freq, days = _frequency(get(r, "frequency"))
        must = _yes(get(r, "must_see"))
        name_norm = norm_name(name)
        inst_norm = norm_text(get(r, "hospital")) or ""
        key = f"e:{email_norm}" if email_norm else (f"p:{phone_norm}" if phone_norm else f"n:{name_norm}|{inst_norm}")

        match = (by_phone.get(phone_norm) if phone_norm else None) or (by_email.get(email_norm) if email_norm else None) \
            or by_name.get((name_norm, inst_norm))
        match_by = None
        if match:
            match_by = "WhatsApp number" if phone_norm and by_phone.get(phone_norm) is match else (
                "Email" if email_norm and by_email.get(email_norm) is match else "Name + hospital")
        if match and dept and match["department"] != dept:
            dept = match["department"]
        dup_of = seen.get(key)
        if dup_of is not None:
            issues.append(f"Duplicate of row {dup_of + 1} in this file")
            summary["duplicates"] += 1
        else:
            seen[key] = len(out)
        row = {
            "row": n + 1, "doctor_name": name, "specialty": specialty or None, "institute": get(r, "hospital") or None,
            "whatsapp_number": wa, "email": email_norm, "department": dept,
            "sub_department": (match or {}).get("sub_department") or sub,
            "city": get(r, "city") or None, "state": get(r, "state") or None, "country": get(r, "country") or None,
            "contact_frequency": freq, "frequency_days": days, "must_see": must,
            "active": get(r, "active") or None, "notes": get(r, "notes") or None,
            "dedupe_key": key, "email_norm": email_norm, "phone_norm": phone_norm, "name_norm": name_norm,
            "existing_doctor_id": match["id"] if match else None, "match_by": match_by,
            "is_new": match is None, "duplicate": dup_of is not None, "issues": issues,
            "blocking": dup_of is not None or not dept or (dept not in allowed_departments) or (not wa and not email_norm),
        }
        areas = doctor_areas({"department": dept, "sub_department": row["sub_department"], "specialty": specialty})
        row["areas"] = areas
        if row["blocking"]:
            summary["errors"] += 1
        summary["existing" if match else "new"] += 0 if dup_of is not None else 1
        summary["oncology"] += "ONCOLOGY" in areas
        summary["hematology"] += "HEMATOLOGY" in areas
        summary["rare_disease"] += "RARE_DISEASE" in areas
        summary["must_see"] += bool(must)
        out.append(row)
    return {"rows": out, "summary": summary}
